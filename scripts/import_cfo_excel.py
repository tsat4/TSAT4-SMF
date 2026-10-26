"""Turn the old TSATP4-CFO workbooks into SQL for Tx Data (tables: db/act/cfo.sql, names mt_cfo_* / tx_cfo_*).

    python -I scripts/import_cfo_excel.py           # reads inbox/master_cfo.xlsx + inbox/transaction_cfo.xlsx
    -> writes inbox/import_cfo.sql (inbox/ is git-ignored: the data stays off GitHub)

Run it in the Tx Data SQL Editor after db/act/cfo.sql. It is safe to run again (on conflict do nothing).
Fixes on the way in (printed at the end): scope 3 EF codes used twice get a suffix, EFs without a value are
Inactive, the septic EF becomes 1 kg CH4 = GWP CH4 kgCO2e. Totals are not imported: v_cfo_entry computes them.
"""
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
GWP = {'gwp-01': 'AR5'}
ROLES = {'ประธาน': 'chairman', 'รองประธาน': 'vice', 'ผู้จัดทำ': 'admin', 'สมาชิก': 'member'}
KIND = {'Mobile Combustion(On-Road)': ('Mobile Combustion', 'On-Road'), 'Mobile Combustion(Off-Road)': ('Mobile Combustion', 'Off-Road')}
COMPANY = {'ชื่อธุรกิจ': 'company_name', 'สถานที่ตั้ง': 'company_address', 'พื้นที่โรงงาน': 'company_area', 'ประเภทของธุรกิจ': 'company_business',
           'กลุ่มลูกค้าหลัก': 'company_customers', 'จำนวนพนักงาน': 'company_employees', 'อัพเดท': 'company_updated', 'รูปภาพ1': 'company_photo1', 'รูปภาพ2': 'company_photo2'}
notes = []


def sheet(book, name):
    rows = [r for r in book[name].iter_rows(values_only=True) if any(c not in (None, '') for c in r)]
    head = [str(c).strip() if c is not None else '' for c in rows[0]]
    return [{k: v for k, v in zip(head, r) if k} for r in rows[1:]]


def lit(value, cast='text'):
    if value is None or (isinstance(value, str) and not value.strip()):
        return f'null::{cast}'
    if isinstance(value, (datetime, date)):
        text = value.strftime('%Y-%m-%d')
    elif isinstance(value, float) and value.is_integer() and cast == 'text':
        text = str(int(value))
    else:
        text = str(value).strip()
    return "'" + text.replace("'", "''") + f"'::{cast}"


def num(value):
    return lit(value if isinstance(value, (int, float)) else None, 'numeric')


def year_of(value):
    digits = ''.join(ch for ch in str(value or '') if ch.isdigit())
    return int(digits[-4:]) if len(digits) >= 4 else None


def values(rows):
    return ',\n  '.join('(' + ', '.join(r) + ')' for r in rows)


def main():
    master = openpyxl.load_workbook(ROOT / 'inbox' / 'master_cfo.xlsx', read_only=True, data_only=True)
    trans = openpyxl.load_workbook(ROOT / 'inbox' / 'transaction_cfo.xlsx', read_only=True, data_only=True)
    out = ['-- CFO data from inbox/*_cfo.xlsx (scripts/import_cfo_excel.py) · run in Tx Data after db/act/cfo.sql · safe to re-run', 'begin;']

    scopes = sheet(master, 'mt_scope')
    code_of = {}
    rows = []
    for s in scopes:
        code = s['scope_id'].replace('scope', '')
        code_of[s['category_eng']] = code
        rows.append([lit(code), lit(int(code.split('-')[0]), 'smallint'), lit(s['category_eng']), lit(s.get('category_th'))])
    out.append(f"insert into mt_cfo_scope (code, scope, name, name_th) values\n  {values(rows)}\non conflict (code) do nothing;")

    gwp = sheet(master, 'mt_gwp')
    rows = [[lit(GWP[g['gwp_id']]), num(g['CO₂']), num(g['Fossil CH₄']), num(g['CH₄']), num(g['N₂O']), num(g['SF₆']), num(g['NF₃']), lit(g.get('Reference')),
             lit('Active' if g.get('status') == 'ACTIVE' else 'Inactive')] for g in gwp]
    out.append(f"insert into mt_cfo_gwp (code, co2, ch4_fossil, ch4, n2o, sf6, nf3, reference, status) values\n  {values(rows)}\non conflict (code) do nothing;")
    gwp_ch4 = gwp[0]['CH₄']

    septic_ef = {a['ef_id'] for a in sheet(master, 'mt_ac_scope1') if a.get('st_id')}
    seen = Counter()
    rows = []
    for name, sc in (('mt_ef_scope1', 1), ('mt_ef_scope2', 2), ('mt_ef_scope3', 3)):
        for e in sheet(master, name):
            cat = e.get('category_eng')
            base, kind = KIND.get(cat, (cat, None))
            if sc == 3:
                kind = str(e['group']).strip() if e.get('group') else None
            scope_code = code_of.get(base)
            if not scope_code and 'น้ำเสีย' in str(kind):
                scope_code = code_of['Waste Generated in Operations']
                notes.append(f"{e['emission_source']}: no category, put in 3-5 Waste Generated in Operations")
            if not scope_code:
                notes.append(f"{e.get('ef_id') or e['emission_source']}: no category, skipped")
                continue
            code = e.get('ef_id')
            if not code:
                # No code in Excel: a stable one from the order, so running the import again adds nothing.
                seen['nocode', sc] += 1
                code = f"sc{sc}-x{seen['nocode', sc]:03d}"
            elif code:
                seen[code] += 1
                if seen[code] > 1:
                    code = f"{code}-{seen[code]}"
                    notes.append(f"{e['ef_id']} used twice in {name}: the later row is {code}")
            total = e.get('kgCO₂e/unit(total)')
            gas = [e.get(k) for k in ('CO₂', 'Fossil CH₄', 'CH₄', 'N₂O', 'HFCₛ', 'PFCₛ')] if sc == 1 else [None] * 6
            if e['ef_id'] in septic_ef:
                gas = [0, 0, 1, 0, 0, 0]
                total = gwp_ch4
                notes.append(f"{e['ef_id']} (septic, 1 kg CH4): kgco2e {e.get('kgCO₂e/unit(total)')} -> {gwp_ch4} = GWP CH4")
            active = e.get('status') == 'ACTIVE' and isinstance(total, (int, float))
            if not isinstance(total, (int, float)):
                notes.append(f"{code or e['emission_source']}: no EF value, Inactive")
            unit = str(e['unit']).strip()
            rows.append([lit(code), lit(scope_code), lit(kind), lit(e['emission_source']), lit(e.get('detail')), lit(unit.lower() if unit in ('Kg', 'Kwh') else unit), num(total),
                         *[num(g) for g in gas], lit(e.get('reference')), lit(year_of(e.get('version')), 'integer'), lit('Active' if active else 'Inactive')])
    out.append("insert into mt_cfo_ef (code, scope_code, kind, name, detail, unit, kgco2e, co2, ch4_fossil, ch4, n2o, hfc, pfc, reference, year, status) values\n  "
               f"{values(rows)}\non conflict (code) do nothing;")

    rows = []
    for name in ('mt_ac_scope1', 'mt_ac_scope2', 'mt_ac_scope3'):
        for a in sheet(master, name):
            rows.append([lit(a['ac_id']), lit(a['ef_id']), lit(GWP[a['gwp_id']]), lit(a['department']), lit(a['activity_name']), lit(a.get('description')),
                         lit('septic' if a.get('st_id') else None), lit('Active' if a.get('status') == 'ACTIVE' else 'Inactive')])
    out.append("insert into mt_cfo_activity (code, ef_id, gwp_code, department, name, detail, calc, status)\n"
               "select v.code, f.id, v.gwp, v.department, v.name, v.detail, v.calc, v.status from (values\n  "
               f"{values(rows)}\n) v(code, ef, gwp, department, name, detail, calc, status)\njoin mt_cfo_ef f on f.code = v.ef\non conflict (code) do nothing;")

    rows = []
    for o in sheet(master, 'mt_org'):
        role = ROLES[o['role']]
        rows.append([lit(o['updated_at'].year if isinstance(o.get('updated_at'), datetime) else 2026, 'integer'), lit(o['employee_id']), lit(role),
                     lit(o['department'] if role == 'member' else None)])
    out.append(f"insert into mt_cfo_team (year, employee_id, role, department) values\n  {values(rows)}\non conflict do nothing;")

    company = sheet(master, 'mt_company')[0]
    rows = [[lit(key), lit(company.get(col))] for col, key in COMPANY.items() if company.get(col) not in (None, '')]
    out.append(f"insert into mt_cfo_setting (key, value) values\n  {values(rows)}\non conflict (key) do nothing;")

    rows = []
    for name in ('tx_scope1', 'tx_scope2', 'tx_scope3'):
        for t in sheet(trans, name):
            src = t.get('source_id')
            rows.append([lit(t['ac_id']), lit(t['month/year'], 'date'), num(t.get('activity_quantity') or 0), lit(src if str(src or '').startswith('sc') else None),
                         lit('CONFIRMED' if t.get('status') == 'CONFIRMED' else 'DRAFT'), lit(t.get('remark'))])
    out.append("insert into tx_cfo_entry (activity_id, month, qty, ef_id, gwp_code, status, remark, created_by)\n"
               "select a.id, v.month, v.qty, coalesce(f.id, a.ef_id), a.gwp_code, v.status, v.remark, 'import' from (values\n  "
               f"{values(rows)}\n) v(ac, month, qty, ef, status, remark)\njoin mt_cfo_activity a on a.code = v.ac\nleft join mt_cfo_ef f on f.code = v.ef\n"
               "on conflict (activity_id, month) do nothing;")
    out.append('commit;')
    out.append("-- check: rows per table and tCO2e per scope (compare with tCO₂e(total) in transaction_cfo.xlsx)\n"
               "select (select count(*) from mt_cfo_ef) ef, (select count(*) from mt_cfo_activity) activities, (select count(*) from mt_cfo_team) team,\n"
               "  (select count(*) from tx_cfo_entry) entries, (select round(sum(tco2e), 3) from v_cfo_entry where scope = 1) scope1_t,\n"
               "  (select round(sum(tco2e), 3) from v_cfo_entry where scope = 2) scope2_t;")
    path = ROOT / 'inbox' / 'import_cfo.sql'
    path.write_text('\n\n'.join(out) + '\n', encoding='utf-8')
    sys.stdout.reconfigure(encoding='utf-8')
    excel = {n: round(sum(t.get('tCO₂e(total)') or 0 for t in sheet(trans, n)), 3) for n in ('tx_scope1', 'tx_scope2')}
    print(f'wrote {path.relative_to(ROOT)} · Excel tCO2e scope1 {excel["tx_scope1"]} scope2 {excel["tx_scope2"]}')
    print('\n'.join(f'  - {n}' for n in notes))


if __name__ == '__main__':
    main()
